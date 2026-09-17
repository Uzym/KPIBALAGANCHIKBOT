"""BLL: экспорт итогов в красивый xlsx (для /экспорт)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dal import repositories as repo
from dal.database import from_db

from config import Config
from .publication import event_vm
from .summary import fmt_when

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws, widths: dict[int, int]) -> None:
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_export_xlsx(s, cfg: Config) -> bytes:
    wb = Workbook()

    # --- События ---
    ws = wb.active
    ws.title = "События"
    headers = ["#id", "Название", "Даты", "Место", "Статус",
               "Придут", "Возможно", "Не придут", "Отмечено пришедших"]
    ws.append(headers)
    for e in sorted(repo.all_events(s), key=lambda x: x.starts_at):
        vm = event_vm(s, cfg, e)
        yes = [x.name for x in vm.lists[0].entries if not x.struck_at]
        maybe = [x.name for x in vm.lists[1].entries if not x.struck_at]
        no = [x.name for x in vm.lists[2].entries if not x.struck_at]
        attended = []
        for att in repo.attendance_for_event(s, e.id):
            if att.person_id is not None:
                vk = repo.vk_account_for_person(s, att.person_id)
                if vk is not None:
                    attended.append(vk.display_name or str(vk.platform_user_id))
            elif att.account_id is not None:
                acc = repo.get_account(s, att.account_id)
                if acc is not None:
                    attended.append(acc.display_name or str(acc.platform_user_id))
        status = {"draft": "черновик", "active": "идёт запись", "closed": "запись закрыта",
                  "cancelled": "отменено", "finished": "завершено"}.get(e.status, e.status)
        ws.append([
            e.id, e.title,
            fmt_when(from_db(e.starts_at), from_db(e.ends_at), cfg.tz),
            e.place or "", status,
            ", ".join(yes), ", ".join(maybe), ", ".join(no), ", ".join(attended),
        ])
    _style_header(ws)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    _autofit(ws, {1: 6, 2: 24, 3: 22, 4: 14, 5: 16, 6: 34, 7: 34, 8: 34, 9: 30})

    # --- Опросы ---
    ws2 = wb.create_sheet("Опросы")
    ws2.append(["#id", "Вопрос", "Срок", "Статус", "Вариант", "Голосов", "Кто выбрал"])
    from . import polls as polls_bll
    for p in sorted(repo.all_polls(s), key=lambda x: x.id):
        vm = polls_bll.poll_vm(s, cfg, p)
        status = {"draft": "черновик", "active": "идёт", "closed": "закрыт",
                  "cancelled": "отменён"}.get(p.status, p.status)
        for o in vm.options:
            current = [e.name for e in o.entries if not e.struck_at]
            ws2.append([
                p.id if o.idx == 0 else "", p.title if o.idx == 0 else "",
                polls_bll.closes_line(p, cfg) or "" if o.idx == 0 else "",
                status if o.idx == 0 else "",
                f"{o.idx + 1}. {o.text}", len(current), ", ".join(current),
            ])
    _style_header(ws2)
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    _autofit(ws2, {1: 6, 2: 26, 3: 18, 4: 12, 5: 22, 6: 9, 7: 40})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
